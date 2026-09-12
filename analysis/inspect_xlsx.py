#!/usr/bin/env python3
"""Inspect all FED_SPA xlsx files: sheets, columns, row counts, sample rows."""
import glob
import openpyxl

files = sorted(glob.glob("/workspace/FED_SPA*.xlsx")) + ["/workspace/FED-SPA_Location_List.xlsx"]
files = sorted(set(files))

for f in files:
    try:
        wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
        print(f"\n{'='*80}")
        print(f"FILE: {f}")
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            print(f"  SHEET: {ws.title!r}  dims={ws.max_row}x{ws.max_column}")
            if rows:
                print(f"    HEADER: {rows[0]}")
                for r in rows[1:4]:
                    print(f"    ROW: {r}")
                print(f"    TOTAL DATA ROWS: {len(rows)-1}")
        wb.close()
    except Exception as e:
        print(f"\nFILE: {f} -> ERROR {e}")
